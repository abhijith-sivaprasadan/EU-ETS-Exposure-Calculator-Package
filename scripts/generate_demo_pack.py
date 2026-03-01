from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import math
import statistics

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "EU_ETS_Exposure_Calculator_Template.xlsx"
OUTPUT_SINGLE = ROOT / "EU_ETS_Exposure_Calculator_Demo_Insights.xlsx"
OUTPUT_PORTFOLIO_DIR = ROOT / "demo_portfolio"
OUTPUT_PORTFOLIO_BOOK = ROOT / "EU_ETS_Portfolio_Demo_Worksheet.xlsx"


@dataclass(frozen=True)
class SiteConfig:
    name: str
    scenario: str
    elec_mult: float
    gas_mult: float
    steam_mult: float
    prod_mult: float
    other_fuel_mwh: float
    event_month_idx: int
    event_scale: float


def month_labels(start_year: int = 2024, n: int = 24) -> list[str]:
    labels: list[str] = []
    year = start_year
    month = 1
    for _ in range(n):
        labels.append(f"{year:04d}-{month:02d}")
        month += 1
        if month > 12:
            month = 1
            year += 1
    return labels


def generate_series(cfg: SiteConfig, n: int = 24) -> list[dict[str, float | str]]:
    rows: list[dict[str, float | str]] = []
    months = month_labels(n=n)
    for i, period in enumerate(months):
        season = math.sin((i / 12) * 2 * math.pi)
        shoulder = math.cos((i / 6) * math.pi)

        elec = (2350 + 150 * season + 70 * shoulder) * cfg.elec_mult
        gas_nm3 = (1_090_000 + 120_000 * season + 55_000 * shoulder) * cfg.gas_mult
        steam = (620 + 65 * season) * cfg.steam_mult
        prod_t = (43_000 + 4_500 * (0.5 + 0.5 * shoulder) + 2_200 * season) * cfg.prod_mult
        elec_price = 59 + 3.2 * shoulder + 1.8 * season

        if i == cfg.event_month_idx:
            elec *= cfg.event_scale
            gas_nm3 *= cfg.event_scale
            steam *= max(0.85, cfg.event_scale - 0.1)
            prod_t *= max(0.88, 1.0 / cfg.event_scale)
            elec_price *= 1.18

        rows.append(
            {
                "period": period,
                "elec_mwh": round(elec, 1),
                "gas_nm3": round(gas_nm3, 0),
                "other_fuel_mwh": round(cfg.other_fuel_mwh if i % 6 in (1, 2) else 0.0, 1),
                "steam_mwh": round(steam, 1),
                "prod_t": round(prod_t, 0),
                "elec_price": round(elec_price, 1),
            }
        )
    return rows


def fill_inputs_sheet(workbook_path: Path, cfg: SiteConfig) -> None:
    wb = load_workbook(workbook_path)
    ws = wb["Inputs"]
    ws["B3"] = cfg.scenario

    for r in range(6, 30):
        for c in range(1, 9):
            ws.cell(r, c).value = None

    rows = generate_series(cfg)
    for i, row in enumerate(rows):
        r = 6 + i
        ws.cell(r, 1).value = row["period"]
        ws.cell(r, 2).value = row["elec_mwh"]
        ws.cell(r, 3).value = row["gas_nm3"]
        ws.cell(r, 4).value = ""
        ws.cell(r, 5).value = row["other_fuel_mwh"]
        ws.cell(r, 6).value = row["steam_mwh"]
        ws.cell(r, 7).value = row["prod_t"]
        ws.cell(r, 8).value = row["elec_price"]

    wb.save(workbook_path)


def build_portfolio_worksheet(site_files: list[Path]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Portfolio"
    ws.append(["Portfolio Demo Worksheet", "", "", "Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.append([])
    ws.append(
        [
            "Workbook",
            "Use in API field",
            "Profile",
            "What to look for",
            "Expected behavior",
        ]
    )

    rows = [
        (
            site_files[0].name,
            "workbooks",
            "Balanced site",
            "Seasonal trend with one stress month",
            "Mid risk under default thresholds",
        ),
        (
            site_files[1].name,
            "workbooks",
            "High gas exposure",
            "Stronger ETS volatility in COST tab",
            "Higher ETS cost and amber/red potential",
        ),
        (
            site_files[2].name,
            "workbooks",
            "Low-carbon relative site",
            "Lower Scope1 and flatter cost curve",
            "Lower risk for same EUA settings",
        ),
    ]
    for row in rows:
        ws.append(row)

    ws.append([])
    ws.append(["Recommended test settings", "", "", "", ""])
    ws.append(["Primary Scenario", "Mid", "", "", ""])
    ws.append(["Base / Stress / High EUA", "85 / 110 / 140", "", "", ""])
    ws.append(["Scope3 factor (tCO2/t)", "0.35", "", "", ""])
    ws.append(["Guardrails (Amber / Red)", "4,000,000 / 6,000,000", "", "", ""])

    ws2 = wb.create_sheet("Scope3_Guide")
    ws2.append(["Scope 3 quick setup guidance"])
    ws2.append([])
    ws2.append(["Variable", "Suggestion", "Notes"])
    ws2.append(["Scope3 factor (tCO2 per t production)", "0.15 to 0.60", "Pick proxy based on upstream intensity"])
    ws2.append(["Initial default for demo", "0.35", "Used in UI for strong visibility of Scope3 trend"])
    ws2.append(["Interpretation", "", "Scope3 is reporting-oriented; ETS cost remains Scope1-only"])

    wb.save(OUTPUT_PORTFOLIO_BOOK)


def make_single_insight_book() -> None:
    OUTPUT_SINGLE.write_bytes(TEMPLATE.read_bytes())
    cfg = SiteConfig(
        name=OUTPUT_SINGLE.name,
        scenario="Mid",
        elec_mult=1.0,
        gas_mult=1.0,
        steam_mult=1.0,
        prod_mult=1.0,
        other_fuel_mwh=120.0,
        event_month_idx=15,
        event_scale=1.22,
    )
    fill_inputs_sheet(OUTPUT_SINGLE, cfg)


def make_portfolio_books() -> list[Path]:
    OUTPUT_PORTFOLIO_DIR.mkdir(parents=True, exist_ok=True)
    configs = [
        SiteConfig(
            name="EU_ETS_Portfolio_Site_A.xlsx",
            scenario="Mid",
            elec_mult=1.00,
            gas_mult=1.00,
            steam_mult=1.00,
            prod_mult=1.00,
            other_fuel_mwh=90.0,
            event_month_idx=8,
            event_scale=1.16,
        ),
        SiteConfig(
            name="EU_ETS_Portfolio_Site_B.xlsx",
            scenario="High",
            elec_mult=1.08,
            gas_mult=1.22,
            steam_mult=1.05,
            prod_mult=0.95,
            other_fuel_mwh=180.0,
            event_month_idx=17,
            event_scale=1.28,
        ),
        SiteConfig(
            name="EU_ETS_Portfolio_Site_C.xlsx",
            scenario="Low",
            elec_mult=0.88,
            gas_mult=0.82,
            steam_mult=0.93,
            prod_mult=1.04,
            other_fuel_mwh=40.0,
            event_month_idx=5,
            event_scale=1.12,
        ),
    ]

    output_files: list[Path] = []
    for cfg in configs:
        out_path = OUTPUT_PORTFOLIO_DIR / cfg.name
        out_path.write_bytes(TEMPLATE.read_bytes())
        fill_inputs_sheet(out_path, cfg)
        output_files.append(out_path)
    return output_files


def summarize_single_book() -> str:
    wb = load_workbook(OUTPUT_SINGLE, data_only=True)
    ws = wb["Inputs"]
    prods = []
    gases = []
    for r in range(6, 30):
        p = ws.cell(r, 7).value
        g = ws.cell(r, 3).value
        if isinstance(p, (int, float)):
            prods.append(float(p))
        if isinstance(g, (int, float)):
            gases.append(float(g))
    if not prods or not gases:
        return "No summary available."
    return (
        f"Prod mean={statistics.mean(prods):,.0f} t, "
        f"Gas mean={statistics.mean(gases):,.0f} Nm3, "
        f"Gas peak={max(gases):,.0f} Nm3"
    )


def main() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")

    make_single_insight_book()
    sites = make_portfolio_books()
    build_portfolio_worksheet(sites)

    print("Generated:")
    print(f"  - {OUTPUT_SINGLE.name}")
    for p in sites:
        print(f"  - {p.relative_to(ROOT)}")
    print(f"  - {OUTPUT_PORTFOLIO_BOOK.name}")
    print("Single-book summary:", summarize_single_book())


if __name__ == "__main__":
    main()
